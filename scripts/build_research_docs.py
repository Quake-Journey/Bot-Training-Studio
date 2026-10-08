"""Build/check self-contained bilingual research DOCX without changing app guides."""
import argparse
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import xml.etree.ElementTree as ET
import zipfile

DOCS=Path(__file__).resolve().parents[1]/'docs'
MANIFEST=DOCS/'q3t2-research-docs.json'


def sha(data):return hashlib.sha256(data).hexdigest()


def validate(data):
    with zipfile.ZipFile(BytesIO(data)) as archive:
        if archive.testzip():raise ValueError('Corrupt research DOCX')
        for name in archive.namelist():
            if not name.endswith(('.xml','.rels')):continue
            raw=archive.read(name);root=ET.fromstring(raw)
            if any(e.get('TargetMode')=='External' for e in root.iter()):raise ValueError('External document relationship')
            if b'instrText' in raw or b'fldSimple' in raw:raise ValueError('Unexpected document field')
            if re.search(rb'[A-Za-z]:[\\/]Claude2|github_pat_|ghp_',raw):raise ValueError('Private content')


def build(language):
    from docx import Document
    from docx.shared import Mm,Pt
    source=DOCS/f'Q3T2_LEARNING_{language}.md';doc=Document()
    section=doc.sections[0];section.page_width=Mm(210);section.page_height=Mm(297)
    section.top_margin=section.bottom_margin=Mm(18)
    section.left_margin=section.right_margin=Mm(20)
    style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(11)
    style.paragraph_format.space_after=Pt(6)
    section.header.paragraphs[0].text='QUAKE JOURNEY / BOT TRAINING STUDIO / RESEARCH'
    lines=source.read_text(encoding='utf8').splitlines();index=0
    def clean(text):return text.replace('**','').replace('`','')
    while index<len(lines):
        line=lines[index].strip();index+=1
        if not line:continue
        if line.startswith('|'):
            rows=[line]
            while index<len(lines) and lines[index].strip().startswith('|'):
                rows.append(lines[index].strip());index+=1
            rows=[r for r in rows if not re.fullmatch(r'[| :\-]+',r)]
            cells=[[clean(v.strip()) for v in r.strip('|').split('|')] for r in rows]
            table=doc.add_table(rows=0,cols=len(cells[0]));table.style='Table Grid'
            for n,row in enumerate(cells):
                for cell,text in zip(table.add_row().cells,row):
                    cell.text=text
                    if n==0:
                        for run in cell.paragraphs[0].runs:run.bold=True
            continue
        if line.startswith('# '):doc.add_heading(clean(line[2:]),0)
        elif line.startswith('## '):doc.add_heading(clean(line[3:]),1)
        elif line.startswith('- '):doc.add_paragraph(clean(line[2:]),style='List Bullet')
        else:
            # Markdown prose wraps are one paragraph, not a line per paragraph.
            parts=[line]
            while index<len(lines) and lines[index].strip() and not lines[index].lstrip().startswith(('#','- ','|')):
                parts.append(lines[index].strip());index+=1
            doc.add_paragraph(clean(' '.join(parts)))
    raw=BytesIO();doc.save(raw);stable=BytesIO()
    with zipfile.ZipFile(raw) as archive,zipfile.ZipFile(stable,'w',zipfile.ZIP_DEFLATED) as output:
        for name in sorted(archive.namelist()):
            info=zipfile.ZipInfo(name,(2026,1,1,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED
            output.writestr(info,archive.read(name))
    data=stable.getvalue();validate(data)
    target=DOCS/f'Bot_Training_Studio_q3t2_Research_{language}.docx';target.write_bytes(data)
    return dict(language=language,source=source.name,source_sha256=sha(source.read_bytes()),file=target.name,sha256=sha(data))


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--check',action='store_true');args=parser.parse_args()
    builder=sha(Path(__file__).read_bytes())
    if args.check:
        meta=json.loads(MANIFEST.read_text(encoding='utf8'))
        if meta['builder_sha256']!=builder or {r['language'] for r in meta['documents']}!={'RU','EN'}:
            raise ValueError('Research DOCX builder/pair changed')
        for row in meta['documents']:
            raw=(DOCS/row['file']).read_bytes();validate(raw)
            if sha(raw)!=row['sha256'] or sha((DOCS/row['source']).read_bytes())!=row['source_sha256']:
                raise ValueError('Stale research DOCX: '+row['file'])
    else:
        MANIFEST.write_text(json.dumps(dict(builder_sha256=builder,documents=[build(lang) for lang in ('RU','EN')]),indent=2)+'\n',encoding='utf8')
    print(json.dumps(dict(research_docs_valid=True)))


if __name__=='__main__':main()
