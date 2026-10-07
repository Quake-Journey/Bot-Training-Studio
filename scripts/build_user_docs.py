"""Build the paired, self-contained DOCX user guides from their tracked Markdown sources."""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import zipfile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
MANIFEST = DOCS / "user-guides.json"
VERSION = ET.parse(ROOT / 'src/BotTrainingStudio/BotTrainingStudio.csproj').findtext('./PropertyGroup/Version')
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def validate(data):
    with zipfile.ZipFile(BytesIO(data)) as archive:
        if archive.testzip() is not None:
            raise ValueError("Corrupt DOCX ZIP")
        texts = []
        for name in archive.namelist():
            raw = archive.read(name)
            if name.endswith((".xml", ".rels")):
                element = ET.fromstring(raw)
                if name.endswith(".rels"):
                    for relation in element:
                        if relation.get("TargetMode") == "External":
                            raise ValueError("External DOCX relationship: " + name)
                fields = [field.text or "" for field in element.findall(".//w:instrText", NS)]
                fields += [field.get("{" + NS["w"] + "}instr", "") for field in element.findall(".//w:fldSimple", NS)]
                for instruction in fields:
                    if instruction.strip() != "PAGE":
                        raise ValueError("Unexpected document field")
                for update in element.findall(".//w:updateFields", NS):
                    if update.get("{" + NS["w"] + "}val", "true") not in ("false", "0", "off"):
                        raise ValueError("Automatic field updates enabled")
                texts.extend(t.text or "" for t in element.findall(".//w:t", NS))
        full = "\n".join(texts)
        if re.search(r"[A-Za-z]:[\\/]Claude2|github_pat_|ghp_", full):
            raise ValueError("Private path or token in documentation")
        if "BotTrainingStudio.exe" not in full or "11.1" in full:
            raise ValueError("Unexpected guide content/version")


def build(language):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    source = DOCS / f"USER_GUIDE_{language}.md"
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(20), Mm(18)
    section.left_margin, section.right_margin = Mm(20), Mm(20)
    section.header_distance = section.footer_distance = Mm(9)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.05
    normal.paragraph_format.widow_control = True
    for name, size in [("Title", 27), ("Heading 1", 18), ("Heading 2", 12)]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("5145A5")
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    doc.styles["Title"].paragraph_format.space_before = Pt(0)
    rpr = normal.element.get_or_add_rPr()
    lang = OxmlElement("w:lang"); lang.set(qn("w:val"), "ru-RU" if language == "RU" else "en-US"); rpr.append(lang)
    settings = doc.settings.element
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "false"); settings.append(update)
    header = section.header.paragraphs[0]
    header.text = "QUAKE JOURNEY   /   BOT TRAINING STUDIO"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("777389")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(f"{language}  •  {VERSION}  •  by ly     |     ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    props = doc.core_properties
    props.title = "Bot Training Studio — " + ("Руководство пользователя" if language == "RU" else "User guide")
    props.author = props.last_modified_by = "Quake Journey"
    props.subject = f"Version {VERSION} development preview"
    props.language = "ru-RU" if language == "RU" else "en-US"
    props.comments = ""
    props.created = props.modified = datetime(2026, 10, 7, tzinfo=timezone.utc)

    def inline(paragraph, text):
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            run = paragraph.add_run(part[2:-2] if part.startswith("**") else part[1:-1] if part.startswith("`") else part)
            if part.startswith("**"): run.bold = True
            if part.startswith("`"): run.font.name = "Consolas"; run.font.size = Pt(9)

    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip(); i += 1
        if not line: continue
        if line == "<!-- page -->": doc.add_page_break(); continue
        if line.startswith("# "):
            doc.add_paragraph(line[2:], "Title"); continue
        if line.startswith("## "):
            doc.add_paragraph(line[3:], "Heading 1"); continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:], "Heading 2"); continue
        if line.startswith("```"):
            while i < len(lines) and not lines[i].startswith("```"):
                p = doc.add_paragraph(); run = p.add_run(lines[i]); run.font.name = "Consolas"; run.font.size = Pt(9)
                i += 1
            i += 1; continue
        if line.startswith("|"):
            rows = [line]
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip()); i += 1
            values = [[v.strip() for v in row.strip("|").split("|")] for row in rows if not re.fullmatch(r"[| :\-]+", row)]
            table = doc.add_table(rows=0, cols=len(values[0])); table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"; table.autofit = False
            table.columns[0].width = Mm(48); table.columns[1].width = Mm(122)
            for n, row in enumerate(values):
                cells = table.add_row().cells
                cells[0].width = Mm(48); cells[1].width = Mm(122)
                trpr = table.rows[-1]._tr.get_or_add_trPr(); trpr.append(OxmlElement("w:cantSplit"))
                if n == 0: trpr.append(OxmlElement("w:tblHeader"))
                for cell, text in zip(cells, row):
                    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(4); p.paragraph_format.space_before = Pt(4)
                    inline(p, text)
                    shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "5145A5" if n == 0 else "F5F4FA" if n % 2 else "FFFFFF")
                    cell._tc.get_or_add_tcPr().append(shade)
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        if n == 0: run.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue
        if line.startswith("- "):
            inline(doc.add_paragraph(style="List Bullet"), line[2:]); continue
        if re.match(r"\d+\. ", line):
            inline(doc.add_paragraph(), line); continue
        inline(doc.add_paragraph(), line)

    raw = BytesIO(); doc.save(raw)
    # Stable ZIP metadata makes regeneration and freshness checks reproducible.
    result = BytesIO()
    with zipfile.ZipFile(raw) as archive, zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as output:
        for name in sorted(archive.namelist()):
            info = zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0)); info.compress_type = zipfile.ZIP_DEFLATED
            output.writestr(info, archive.read(name))
    data = result.getvalue(); validate(data)
    target = DOCS / f"Bot_Training_Studio_User_Guide_{language}.docx"
    target.write_bytes(data)
    return dict(language=language, source=source.name, source_sha256=sha(source.read_bytes()), file=target.name, sha256=sha(data))


def check():
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    if manifest['version'] != VERSION: raise ValueError('Guide version does not match the application')
    if manifest["builder_sha256"] != sha(Path(__file__).read_bytes()): raise ValueError("DOCX builder changed; rebuild both guides")
    if {entry["language"] for entry in manifest["guides"]} != {"RU", "EN"}: raise ValueError("Both guides are required")
    for entry in manifest["guides"]:
        if sha((DOCS / entry["source"]).read_bytes()) != entry["source_sha256"]: raise ValueError("DOCX source changed: " + entry["source"])
        raw = (DOCS / entry["file"]).read_bytes()
        if sha(raw) != entry["sha256"]: raise ValueError("DOCX file changed: " + entry["file"])
        validate(raw)


def main():
    parser = argparse.ArgumentParser(); parser.add_argument("--check", action="store_true"); args = parser.parse_args()
    if args.check: check()
    else:
        guides = [build(language) for language in ("RU", "EN")]
        MANIFEST.write_text(json.dumps(dict(schema=1, version=VERSION, builder_sha256=sha(Path(__file__).read_bytes()), guides=guides), indent=2) + "\n", encoding="utf-8")
    print(json.dumps(dict(pass_=True, guides=2, mode="check" if args.check else "build")))


if __name__ == "__main__": main()
